"""
resume_parser.py — GDMR Connect
=================================
Turn an uploaded resume (PDF / DOCX) into candidate-form suggestions.

Two layers:
  1. extract_text()   — pure-Python text extraction (pypdf, python-docx).
  2. heuristics()     — regex/keyword extraction that needs no network: email,
                        phone, name, skills, education, experience, location,
                        notice period. Contact details are ALWAYS taken from
                        here, never from an LLM.

routes/ats.py optionally layers an LLM pass on top for the fields heuristics are
weak at (current company, job title, a cleaner skills list); contact details
are masked out of the text before that call (see mask_contacts).

Everything returned is a *suggestion* — the Add Candidate form highlights
auto-filled fields and makes the recruiter confirm them before saving.
"""
import io
import re

MAX_PAGES = 12
MAX_CHARS = 40000

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
# Indian mobile: optional +91/91/0, then 10 digits starting 6-9, allowing a space/dash in the middle
PHONE_RE = re.compile(r"(?<!\d)(?:\+?91[\s-]?|0)?([6-9]\d{4}[\s-]?\d{5})(?!\d)")
URL_RE   = re.compile(r"(?:https?://|www\.)\S+|linkedin\.com/\S+|github\.com/\S+", re.I)

SECTION_WORDS = (
    "skills", "technical skills", "key skills", "core competencies", "skill set", "expertise", "tools",
    "experience", "work experience", "professional experience", "employment", "education",
    "educational qualification", "qualifications", "academic", "projects", "certifications",
    "achievements", "summary", "profile", "objective", "career objective", "declaration",
    "personal details", "personal information", "languages", "interests", "hobbies", "references",
    "internship", "internships", "training", "courses",
)

# (canonical label, regex) — ordered highest qualification first
DEGREES = (
    ("PhD",         r"\bph\.?\s?d\b|\bdoctorate\b"),
    ("MBA",         r"\bm\.?\s?b\.?\s?a\b|\bpgdm\b"),
    ("MCA",         r"\bm\.?\s?c\.?\s?a\b"),
    ("M.Tech",      r"\bm\.?\s?tech\b|\bm\.?\s?e\b(?=\s*[\(\-,])"),
    ("M.Sc",        r"\bm\.?\s?sc\b"),
    ("M.Com",       r"\bm\.?\s?com\b"),
    ("MSW",         r"\bm\.?\s?s\.?\s?w\b"),
    ("MA",          r"\bm\.?\s?a\b(?=\s*[\(\-,]|\s+in\b|\s+english|\s+economics|\s+sociology)|master of arts"),
    ("Master's",    r"\bmaster'?s?\b|\bpost[\s-]?graduat"),
    ("B.Tech",      r"\bb\.?\s?tech\b|\bb\.?\s?e\b(?=\s*[\(\-,]|\s+in\b)|bachelor of engineering"),
    ("BCA",         r"\bb\.?\s?c\.?\s?a\b"),
    ("BBA",         r"\bb\.?\s?b\.?\s?a\b"),
    ("B.Com",       r"\bb\.?\s?com\b|bachelor of commerce"),
    ("B.Sc",        r"\bb\.?\s?sc\b|bachelor of science"),
    ("BSW",         r"\bb\.?\s?s\.?\s?w\b"),
    ("BA",          r"\bb\.?\s?a\b(?=\s*[\(\-,]|\s+in\b|\s+english|\s+economics|\s+sociology)|bachelor of arts"),
    ("Bachelor's",  r"\bbachelor'?s?\b|\bgraduat(?:e|ion)\b|\bundergraduate\b"),
    ("Diploma",     r"\bdiploma\b"),
    ("Plus Two",    r"\bplus\s?two\b|\bhigher secondary\b|\b12th\b|\bxii\b|\bhsc\b"),
    ("10th",        r"\b10th\b|\bsslc\b|\bssc\b|\bmatriculation\b"),
)

CITIES = (
    "Thiruvananthapuram", "Trivandrum", "Kochi", "Cochin", "Ernakulam", "Kozhikode", "Calicut", "Thrissur",
    "Kollam", "Alappuzha", "Alleppey", "Kannur", "Kottayam", "Palakkad", "Malappuram", "Pathanamthitta",
    "Idukki", "Wayanad", "Kasaragod", "Bengaluru", "Bangalore", "Chennai", "Coimbatore", "Mumbai", "Pune",
    "Delhi", "New Delhi", "Noida", "Gurgaon", "Gurugram", "Hyderabad", "Kolkata", "Ahmedabad", "Mangalore",
    "Mysuru", "Mysore", "Madurai", "Remote",
)

# Common skills to spot when there's no explicit "Skills" section.
SKILL_BANK = (
    "Python", "Java", "JavaScript", "TypeScript", "React", "Node.js", "Angular", "Vue", "HTML", "CSS", "SQL",
    "MySQL", "MongoDB", "PostgreSQL", "Excel", "Advanced Excel", "Power BI", "Tableau", "Tally", "SAP", "Photoshop",
    "Illustrator", "After Effects", "Premiere Pro", "Figma", "Canva", "Video Editing", "Motion Graphics",
    "Graphic Design", "UI/UX", "Social Media Marketing", "SEO", "Content Writing", "Copywriting",
    "Digital Marketing", "Recruitment", "Talent Acquisition", "Resume Screening", "Interviewing", "Onboarding",
    "Payroll", "HR Operations", "Communication", "Negotiation", "Project Management", "Team Management",
    "Customer Service", "Sales", "Lead Generation", "MS Office", "Word", "PowerPoint", "Documentation",
    "Data Entry", "Accounting", "GST", "Public Speaking", "Leadership",
)


# ── Text extraction ──────────────────────────────────────────────────────────

def extract_text(data: bytes, filename: str) -> str:
    """Best-effort text from a PDF or DOCX. Returns '' when nothing readable
    (scanned image PDFs, legacy .doc, corrupt files) — callers treat that as
    'ask the recruiter to fill the form by hand'."""
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf") or data[:5] == b"%PDF-":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            parts = []
            for page in reader.pages[:MAX_PAGES]:
                parts.append(page.extract_text() or "")
            return "\n".join(parts)[:MAX_CHARS]
        if name.endswith(".docx"):
            from docx import Document
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:                       # many CV templates are tables
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            return "\n".join(parts)[:MAX_CHARS]
    except Exception as exc:
        print(f"[resume] text extraction failed for {filename!r}: {exc}")
    return ""


def mask_contacts(text: str) -> str:
    """Strip emails / phone numbers / URLs before text leaves the server."""
    text = EMAIL_RE.sub("[email]", text)
    text = PHONE_RE.sub("[phone]", text)
    return URL_RE.sub("[link]", text)


# ── Heuristics ───────────────────────────────────────────────────────────────

def _lines(text):
    return [re.sub(r"\s+", " ", ln).strip() for ln in text.splitlines() if ln.strip()]


def _is_heading(line):
    low = line.lower().strip(" :-–|")
    return low in SECTION_WORDS or (len(line) < 32 and line.isupper() and low.split(" ")[0] in SECTION_WORDS)


def find_email(text):
    m = EMAIL_RE.search(text)
    return m.group(0).lower() if m else ""


def find_phone(text):
    m = PHONE_RE.search(text)
    return re.sub(r"\D", "", m.group(1)) if m else ""


def find_name(lines, email=""):
    skip = re.compile(r"resume|curriculum|vitae|\bcv\b|profile|contact|address|objective|summary|@|\d|http|www\.", re.I)
    for ln in lines[:8]:
        clean = re.sub(r"[^A-Za-z .'-]", " ", ln).strip()
        words = clean.split()
        if skip.search(ln) or not (1 < len(words) <= 4) or _is_heading(ln):
            continue
        if all(len(w) <= 1 for w in words):
            continue
        name = " ".join(words)
        return name.title() if name.isupper() or name.islower() else name
    return ""


def find_skills(text, lines):
    skills, in_section = [], False
    for ln in lines:
        low = ln.lower().strip(" :-–|")
        head = re.match(r"^(?:key |technical |core |it )?(?:skills?(?: set)?|competenc(?:y|ies)|expertise|tools(?: & technologies)?)\b\s*[:\-–]?\s*(.*)$", ln, re.I)
        if head:
            in_section = True
            rest = head.group(1)
            if rest:
                skills += re.split(r"[,;|•·●▪◦•]", rest)
            continue
        if in_section:
            if _is_heading(ln):
                break
            skills += re.split(r"[,;|•·●▪◦•]", ln.lstrip("-–*• "))
            if len(skills) > 40:
                break
    cleaned = []
    for s in skills:
        s = re.sub(r"\s+", " ", s).strip(" .:-–")
        if 1 < len(s) <= 32 and not s.lower().startswith(("and ", "etc")) and s.lower() not in {c.lower() for c in cleaned}:
            cleaned.append(s)
    if len(cleaned) < 3:   # no usable section — spot well-known skills anywhere in the text
        for sk in SKILL_BANK:
            if re.search(r"(?<![A-Za-z])" + re.escape(sk) + r"(?![A-Za-z])", text, re.I) and sk.lower() not in {c.lower() for c in cleaned}:
                cleaned.append(sk)
    return cleaned[:15]


def find_education(text):
    for label, pattern in DEGREES:
        if re.search(pattern, text, re.I):
            return label
    return ""


def find_experience(text):
    """Years of experience as a string ('' when unknown)."""
    best = None
    for m in re.finditer(r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)\b(?:\s+of)?(?:\s+(?:total|work|professional|relevant|industry))?\s*(?:experience)?", text, re.I):
        ctx = text[max(0, m.start() - 40): m.end() + 40].lower()
        if "experience" in ctx or "exp" in ctx:
            val = float(m.group(1))
            if 0 < val <= 45:
                best = max(best or 0, val)
    if best is not None:
        return str(int(best)) if best == int(best) else str(best)
    if re.search(r"\bfresher\b|\bfresh graduate\b|\bno (?:prior )?(?:work )?experience\b", text, re.I):
        return "0"
    return ""


def find_location(text):
    m = re.search(r"(?:current\s+)?(?:location|city|address)\s*[:\-–]\s*([^\n|]{2,60})", text, re.I)
    if m:
        val = re.sub(r"\s+", " ", m.group(1)).strip(" ,.-")
        for city in CITIES:
            if re.search(r"\b" + re.escape(city) + r"\b", val, re.I):
                return city
    for city in CITIES:
        if re.search(r"\b" + re.escape(city) + r"\b", text[:1500], re.I):
            return city
    return ""


def find_notice_period(text):
    m = re.search(r"notice\s+period\s*[:\-–]?\s*([^\n|,;]{1,30})", text, re.I)
    if m:
        val = re.sub(r"\s+", " ", m.group(1)).strip(" .:-")
        if re.search(r"\d|immediate|serving|nil|none", val, re.I):
            return val[:30]
    if re.search(r"\bimmediate(?:ly)?\s+(?:joiner|joining|available)\b", text, re.I):
        return "Immediate"
    return ""


def find_company(lines):
    for ln in lines:
        m = re.match(r"^(?:current|present|last)\s+(?:company|employer|organi[sz]ation)\s*[:\-–]\s*(.{2,60})$", ln, re.I)
        if m:
            return m.group(1).strip(" ,.")
    return ""


def heuristics(text: str) -> dict:
    lines = _lines(text)
    email = find_email(text)
    return {
        "name":             find_name(lines, email),
        "email":            email,
        "phone":            find_phone(text),
        "skills":           find_skills(text, lines),
        "education":        find_education(text),
        "experience":       find_experience(text),
        "current_company":  find_company(lines),
        "current_location": find_location(text),
        "notice_period":    find_notice_period(text),
        "job_title":        "",
    }
